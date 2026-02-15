import os
import sys
import logging
import traceback
import tempfile
from datetime import datetime

# --- 0. EARLY DEBUG LOGGING (Run before anything else) ---
try:
    with open("startup_debug.txt", "w") as f:
        f.write(f"Startup at {datetime.now()}\n")
        f.write(f"Executable: {sys.executable}\n")
        f.write(f"CWD: {os.getcwd()}\n")
        f.write(f"Path: {sys.path}\n")
except Exception as e:
    pass # If we can't write, we can't write.

# --- 1. Setup Logging FIRST (Before any other imports) ---
def setup_logging():
    log_filename = "batch_image_editor_debug.log"
    # Prioritize execution directory (where the exe is)
    exe_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
    current_dir = os.getcwd()
    
    paths_to_try = [
        # 1. Executable directory (User requested "same folder as exe")
        os.path.join(exe_dir, log_filename),
        # 2. Current Working Directory (Often the same, but good failover)
        os.path.join(current_dir, log_filename),
        # 3. Documents (Fallback)
        os.path.join(os.path.expanduser("~"), "Documents", log_filename),
        # 4. Temp (Last resort)
        os.path.join(tempfile.gettempdir(), log_filename),
    ]
    
    for p in paths_to_try:
        try:
            with open(p, 'a') as f:
                f.write(f"\n--- Log Init: {datetime.now()} ---\n")
            return p
        except Exception:
            continue
    return None

log_file = setup_logging()
if log_file:
    logging.basicConfig(
        filename=log_file,
        level=logging.DEBUG,
        format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
    )
else:
    logging.basicConfig(level=logging.DEBUG)

logging.info(f"Backend starting... Executable: {sys.executable}")
logging.info("VERSION: 2.1 - Corrected Filename Logic (FIXED)")
logging.info(f"CWD: {os.getcwd()}")
logging.info(f"Path: {sys.path}")

# --- 2. Robust Imports ---
try:
    from flask import Flask, jsonify, request
    from flask_cors import CORS
    from PIL import Image, ImageOps
    import img2pdf
except ImportError as e:
    logging.critical(f"CRITICAL IMPORT ERROR: {e}")
    logging.critical(traceback.format_exc())
    # If missing DLLs for these, we must exit, but at least we logged it.
    sys.exit(1)
except Exception as e:
    logging.critical(f"CRITICAL STARTUP ERROR: {e}")
    logging.critical(traceback.format_exc())
    sys.exit(1)

app = Flask(__name__)
# Allow CORS for all domains for now, or restrictive for production
CORS(app)

import img2pdf
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError as e:
    logging.critical(f"Failed to import pptx: {e}")
    logging.critical(traceback.format_exc())

@app.route('/health')
def health_check():
    logging.info("Health check requested")
    return jsonify({"status": "ok", "message": "Backend is running"})

import shutil
import tempfile

# Helper function for image transformations
def apply_image_transforms(img, data):
    target_format = data.get('format') or data.get('output_format') # e.g., 'PNG', 'JPEG', 'WEBP'
    resize_target = data.get('resize') # e.g., [800, 600] or 800 (width, maintain aspect)
    aspect_str = data.get('aspect_ratio', 'Original') # '16:9', '4:3', 'Original'
    resize_mode = data.get('resize_mode', 'Contain')   # 'Contain', 'Cover', 'Stretch'

    # --- Aspect Ratio / Resize Logic ---
    if aspect_str != 'Original':
        # Parse aspect ratio
        if aspect_str == '16:9':
            target_ratio = 16 / 9
        elif aspect_str == '4:3':
            target_ratio = 4 / 3
        elif aspect_str == '1:1':
            target_ratio = 1
        else:
            target_ratio = None # Fallback
        
        if target_ratio:
            # Determine target dimensions
            # Use provided width if available, else original width
            t_width = img.width
            if resize_target and isinstance(resize_target, int):
                t_width = resize_target
            
            t_height = int(t_width / target_ratio)
            target_size = (t_width, t_height)

            # Get crop position if available
            centering = (0.5, 0.5) # Default to center
            if 'crop_positions' in data and hasattr(img, 'filename'):
                 # Need to match filename/path. The dictionary keys are full paths.
                 # img.filename usually stores the path if opened from file.
                 
                 positions = data.get('crop_positions', {})
                 # Try exact match first
                 if img.filename in positions:
                     pos = positions[img.filename]
                     centering = (pos['x'] / 100.0, pos['y'] / 100.0)
                 else:
                     # Try normalizing slashes of img.filename to match frontend
                     norm_path = img.filename.replace('\\', '/')
                     if norm_path in positions:
                         pos = positions[norm_path]
                         centering = (pos['x'] / 100.0, pos['y'] / 100.0)

            # Apply resize mode
            if resize_mode == 'Cover':
                # Crop to fill
                img = ImageOps.fit(img, target_size, method=Image.Resampling.LANCZOS, centering=centering)
            elif resize_mode == 'Stretch':
                img = img.resize(target_size, resample=Image.Resampling.LANCZOS)
            else:
                # Default to Contain for 'Contain' or 'Original' or unknown
                # Fit inside and pad
                img = ImageOps.pad(img, target_size, method=Image.Resampling.LANCZOS, color='black', centering=(0.5, 0.5))
            
    elif resize_target:
        # Original aspect ratio but resized
        if isinstance(resize_target, list) and len(resize_target) == 2:
            img = img.resize((resize_target[0], resize_target[1]))
        elif isinstance(resize_target, int):
            # Resize by width, maintain aspect ratio
            w_percent = (resize_target / float(img.size[0]))
            h_size = int((float(img.size[1]) * float(w_percent)))
            img = img.resize((resize_target, h_size))
            
    return img

@app.route('/process', methods=['POST'])
def process_images():
    data = request.json
    print(f"DEBUG: Received data: {data}") # Log the entire payload
    if not data or 'images' not in data:
        return jsonify({"error": "No images provided"}), 400

    images = data['images']
    target_format = data.get('format') or data.get('output_format')

    output_dir = data.get('output_dir')
    overwrite = data.get('overwrite', False)

    # --- Pre-check for existing files ---
    if not overwrite:
        existing_files = []
        for i, img_path in enumerate(images):
             with Image.open(img_path) as img:
                 original_format = img.format
                 save_format = target_format if target_format else original_format
                 if not save_format: save_format = 'PNG'
                 
                 directory, filename = os.path.split(img_path)
                 if output_dir: directory = output_dir
                 name, ext = os.path.splitext(filename)
                 
                 # Construct 3-digit index prefix
                 index_prefix = f"{i+1:03d}"
                 
                 new_filename = f"{index_prefix}_{name}.{save_format.lower()}"
                 new_path = os.path.join(directory, new_filename)
                 
                 if os.path.exists(new_path):
                     existing_files.append(new_filename)
        
        if existing_files:
            return jsonify({"status": "conflict", "existing_files": existing_files, "message": "Files already exist"}), 409

    results = []

    for i, img_path in enumerate(images):
        try:
            with Image.open(img_path) as img:
                original_format = img.format
                
                # Apply transforms
                img = apply_image_transforms(img, data)

                # --- Format Conversion ---
                save_format = target_format if target_format else original_format
                if not save_format:
                     save_format = 'PNG' # Default fallback
                
                # Construct new filename
                directory, filename = os.path.split(img_path)
                
                # Use output_dir if provided, else use original directory
                if output_dir:
                    directory = output_dir
                
                name, ext = os.path.splitext(filename)
                # Construct 3-digit index prefix
                index_prefix = f"{i+1:03d}"
                
                new_filename = f"{index_prefix}_{name}.{save_format.lower()}"
                new_path = os.path.join(directory, new_filename)
                print(f"DEBUG: Saving to {new_path}")

                # Convert to RGB if saving as JPEG (and original was RGBA/P)
                if save_format.upper() == 'JPEG' and img.mode in ('RGBA', 'P'):
                    img = img.convert('RGB')

                img.save(new_path, format=save_format)
                results.append({"original": img_path, "processed": new_path, "status": "success"})
        except Exception as e:
            results.append({"original": img_path, "error": str(e), "status": "error"})

    return jsonify({"results": results})

@app.route('/convert-pdf', methods=['POST'])
def convert_to_pdf():
    data = request.json
    if not data or 'images' not in data:
        return jsonify({"error": "No images provided"}), 400

    images_list = data['images']
    if not images_list:
        return jsonify({"error": "Empty image list"}), 400

    output_dir = data.get('output_dir')
    overwrite = data.get('overwrite', False)

    # Pre-check PDF path
    first_image = images_list[0]
    directory, filename = os.path.split(first_image)
    if output_dir: directory = output_dir
    name, _ = os.path.splitext(filename)
    output_path = os.path.join(directory, f"{name}_combined.pdf")

    if not overwrite and os.path.exists(output_path):
        return jsonify({"status": "conflict", "existing_files": [os.path.basename(output_path)], "message": "File already exists"}), 409

    temp_dir = tempfile.mkdtemp()
    try:
        # Process images first (apply resize/aspect ratio)
        processed_images = []
        for i, img_path in enumerate(images_list):
            try:
                with Image.open(img_path) as img:
                    img = apply_image_transforms(img, data)
                    
                    # Save to temp dir
                    # Convert to RGB to ensure compatibility with various output formats if needed, 
                    # but usually PNG/JPEG is fine.
                    if img.mode == 'RGBA':
                        # Convert to RGB if strictly needed (e.g. for JPEG), 
                        # but PDF supports transparency. 
                        # However, let's keep it simple and ensure no compatibility issues.
                        pass

                    temp_path = os.path.join(temp_dir, f"temp_{i}.jpg")
                    # Save as JPEG for compression or PNG for quality? 
                    # Let's use JPEG for PDF to keep size reasonable, or respect original?
                    # For now, let's force convert to RGB and save as JPEG standard for documents
                    if img.mode != 'RGB':
                        img = img.convert('RGB')
                    img.save(temp_path, quality=90)
                    processed_images.append(temp_path)
            except Exception as e:
                print(f"Skipping image {img_path} due to error: {e}")

        if not processed_images:
             return jsonify({"error": "No valid images to process"}), 400

        # Define output path
        output_dir = data.get('output_dir')
        
        first_image = images_list[0]
        directory, filename = os.path.split(first_image)
        
        if output_dir:
            directory = output_dir

        name, _ = os.path.splitext(filename)
        output_path = os.path.join(directory, f"{name}_combined.pdf")

        with open(output_path, "wb") as f:
            f.write(img2pdf.convert(processed_images))
            
        return jsonify({"status": "success", "processed": output_path, "count": len(processed_images)})

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        shutil.rmtree(temp_dir)

@app.route('/convert-pptx', methods=['POST'])
def convert_to_pptx():
    data = request.json
    if not data or 'images' not in data:
        return jsonify({"error": "No images provided"}), 400

    images_list = data['images']
    if not images_list:
        return jsonify({"error": "Empty image list"}), 400

    temp_dir = tempfile.mkdtemp()
    try:
        # Process images first
        processed_images = []
        max_width_px = 0
        max_height_px = 0

        for i, img_path in enumerate(images_list):
            try:
                with Image.open(img_path) as img:
                    img = apply_image_transforms(img, data)
                    
                    max_width_px = max(max_width_px, img.width)
                    max_height_px = max(max_height_px, img.height)

                    temp_path = os.path.join(temp_dir, f"temp_{i}.png")
                    img.save(temp_path)
                    processed_images.append(temp_path)
            except Exception as e:
                print(f"Skipping image {img_path}: {e}")

        if not processed_images:
            return jsonify({"error": "No valid images found for PPTX conversion"}), 400

        prs = Presentation()
        
        EMU_PER_PX = 9525
        
        if max_width_px == 0 or max_height_px == 0:
             max_width_px = 1280
             max_height_px = 720

        prs.slide_width = max_width_px * EMU_PER_PX
        prs.slide_height = max_height_px * EMU_PER_PX

        blank_slide_layout = prs.slide_layouts[6] 

        first_image = images_list[0]
        directory, filename = os.path.split(first_image)
        name, _ = os.path.splitext(filename)
        output_path = os.path.join(directory, f"{name}_presentation.pptx")

        for img_path in processed_images:
            slide = prs.slides.add_slide(blank_slide_layout)
            
            # Since we already resized images (if settings applied), we might just want to center them
            # We need to know the size of the *processed* image to center it
            with Image.open(img_path) as im:
                 w, h = im.size
            
            pic = slide.shapes.add_picture(img_path, 0, 0)
            
            # Center the image
            pic.left = int((prs.slide_width - pic.width) / 2)
            pic.top = int((prs.slide_height - pic.height) / 2)

        prs.save(output_path)
        return jsonify({"status": "success", "processed": output_path, "count": len(processed_images)})

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        shutil.rmtree(temp_dir)

import fitz  # PyMuPDF

@app.route('/extract-pdf-pages', methods=['POST'])
def extract_pdf_pages():
    data = request.json
    if not data or 'pdf_path' not in data:
        return jsonify({"error": "No pdf_path provided"}), 400

    pdf_path = data['pdf_path']
    if not os.path.exists(pdf_path):
        return jsonify({"error": "File not found"}), 404

    try:
        doc = fitz.open(pdf_path)
        extracted_images = []
        
        # Create a directory for extracted images if it doesn't exist
        # We'll use a specific folder in the same directory as the PDF
        directory, filename = os.path.split(pdf_path)
        name, _ = os.path.splitext(filename)
        output_dir = os.path.join(directory, f"{name}_pages")
        os.makedirs(output_dir, exist_ok=True)

        for i, page in enumerate(doc):
            # Render page to image (pixmap)
            # matrix=fitz.Matrix(2, 2) makes it 2x zoom -> higher resolution (approx 144 dpi)
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2)) 
            image_filename = f"page_{i+1:03d}.png"
            image_path = os.path.join(output_dir, image_filename)
            pix.save(image_path)
            extracted_images.append(image_path)
            
        doc.close()
        return jsonify({"status": "success", "images": extracted_images})

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    # Dynamic Port Selection with Retry Logic
    import socket
    import random
    
    def get_free_port():
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.bind(('', 0))
            return s.getsockname()[1]

    # Check if PORT env var is set
    env_port = os.environ.get("PORT")
    
    MAX_RETRIES = 3
    
    # Disable debug mode in production to avoid reloader
    is_frozen = getattr(sys, 'frozen', False)
    
    for attempt in range(MAX_RETRIES + 1):
        try:
            if env_port:
                port = int(env_port)
            else:
                port = get_free_port()

            logging.info(f"Attempting to start on port: {port} (Attempt {attempt+1}/{MAX_RETRIES+1})")
            
            # CRITICAL: Print port for Electron to capture
            # We print BEFORE run, assuming run will succeed.
            # If run fails, we loop and print again. Electron should handle seeing multiple PORT lines 
            # or just take the valid one. 
            # Note: Electron main.ts logic captures the first one? 
            # Let's check main.ts -> it listens to 'data'. If we print multiple, it might get confused 
            # or just update. 
            # We should probably print *after* successful bind? 
            # But Flask run() blocks. We can't print after.
            # We print *before*. If it fails, we print a new one.
            print(f"PORT:{port}", flush=True)

            # Write port to file
            try:
                with open("backend.port", "w") as f:
                    f.write(str(port))
            except Exception:
                pass

            # Attempt to run
            # Note: 127.0.0.1 is cleaner, but localhost is usually fine.
            app.run(host='127.0.0.1', port=port, debug=not is_frozen)
            
            # If app.run returns, it means server stopped gracefully.
            break
            
        except OSError as e:
            logging.error(f"Port {port} failed: {e}")
            if env_port:
                logging.error("Specified PORT environment variable failed. Cannot retry.")
                break
            
            if attempt < MAX_RETRIES:
                logging.info("Retrying with new port...")
                continue
            else:
                logging.critical("Max retries reached. Exiting.")
                sys.exit(1)
        except Exception as e:
            logging.critical(f"Unexpected error starting server: {e}")
            sys.exit(1)
